"""Native metadata writers per format. mtime is preserved."""

from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from automafile.log import get_logger
from automafile.metadata.mtime import preserve_times


log = get_logger(__name__)


XMP_NAMESPACE = "http://automafile.local/schema/1.0/"
XMP_PREFIX = "automafile"


class NativeMetadataError(Exception):
    """Raised when a native writer cannot complete; caller should fall back."""


def _utc_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")


def _trim(s: str | None, n: int) -> str | None:
    if s is None:
        return None
    s = str(s)
    return s[:n] if len(s) > n else s


def supports(file_path: Path) -> bool:
    return file_path.suffix.lower() in {
        ".pdf", ".docx", ".xlsx", ".pptx",
        ".jpg", ".jpeg", ".png", ".tiff", ".tif", ".webp",
    }


def write(file_path: Path, fields: dict[str, Any]) -> None:
    """Write metadata into the file in place. Raises ``NativeMetadataError`` on failure."""
    suffix = file_path.suffix.lower()
    log.debug("native write: %s (%s)", file_path.name, suffix)
    try:
        with preserve_times(file_path):
            if suffix == ".pdf":
                _write_pdf(file_path, fields)
            elif suffix == ".docx":
                _write_docx(file_path, fields)
            elif suffix == ".xlsx":
                _write_xlsx(file_path, fields)
            elif suffix == ".pptx":
                _write_pptx(file_path, fields)
            elif suffix in {".jpg", ".jpeg", ".png", ".tiff", ".tif", ".webp"}:
                _write_image(file_path, fields)
            else:
                raise NativeMetadataError(f"No native writer for {suffix}")
    except NativeMetadataError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise NativeMetadataError(f"Native write failed for {file_path}: {exc}") from exc


def _write_pdf(file_path: Path, fields: dict[str, Any]) -> None:
    try:
        import pikepdf
    except ImportError as exc:
        raise NativeMetadataError("pikepdf is not installed") from exc

    title = _trim(fields.get("title"), 200)
    summary = _trim(fields.get("summary"), 500)
    tags = fields.get("tags") or []
    keywords = ", ".join(str(t) for t in tags) if tags else None
    category = fields.get("category")
    correspondent = fields.get("correspondent")
    document_date = fields.get("date")
    confidence = fields.get("confidence")

    try:
        pdf = pikepdf.open(file_path, allow_overwriting_input=True)
    except pikepdf.PasswordError as exc:
        raise NativeMetadataError("PDF is encrypted") from exc

    try:
        with pdf.open_metadata() as xmp:
            try:
                xmp.register_xml_namespace(XMP_NAMESPACE, XMP_PREFIX)
            except Exception:
                pass
            if title is not None:
                xmp["dc:title"] = title
            if summary is not None:
                xmp["dc:description"] = summary
            if tags:
                xmp["dc:subject"] = list(tags)
            if category:
                xmp[f"{XMP_PREFIX}:Category"] = str(category)
            if correspondent:
                xmp[f"{XMP_PREFIX}:Correspondent"] = str(correspondent)
            if document_date:
                xmp[f"{XMP_PREFIX}:DocumentDate"] = str(document_date)
            if confidence:
                xmp[f"{XMP_PREFIX}:Confidence"] = str(confidence)

        if pdf.docinfo is None:
            pdf.docinfo = pikepdf.Dictionary()
        if title is not None:
            pdf.docinfo["/Title"] = title
        if summary is not None:
            pdf.docinfo["/Subject"] = summary
        if keywords:
            pdf.docinfo["/Keywords"] = keywords

        pdf.save(file_path)
    finally:
        pdf.close()


def _set_or_create_custom_property(props, name: str, value: Any) -> None:
    """python-docx custom properties helper. Falls back to silent if unsupported."""
    try:
        if name in props:
            props[name].value = value
        else:
            props[name] = value
    except Exception as exc:
        log.debug("Custom prop %s could not be set: %s", name, exc)


def _write_office_core(doc, fields: dict[str, Any]) -> None:
    cp = doc.core_properties
    title = _trim(fields.get("title"), 200)
    summary = _trim(fields.get("summary"), 500)
    tags = fields.get("tags") or []
    keywords = ", ".join(str(t) for t in tags) if tags else None
    category = fields.get("category")
    if title is not None:
        cp.title = title
    if summary is not None:
        cp.subject = summary
        cp.comments = summary
    if keywords:
        cp.keywords = keywords
    if category:
        cp.category = str(category)


def _write_office_custom(doc, fields: dict[str, Any]) -> None:
    """Best-effort write of Automafile-specific keys to custom properties."""
    try:
        from docx.opc.constants import CONTENT_TYPE
    except Exception:
        CONTENT_TYPE = None  # type: ignore[assignment]
    payload = {
        "AutomafileCategory": fields.get("category"),
        "AutomafileCorrespondent": fields.get("correspondent"),
        "AutomafileDocumentDate": fields.get("date"),
        "AutomafileConfidence": fields.get("confidence"),
        "AutomafileMetadataModified": _utc_iso(),
        "AutomafileTags": json.dumps(fields.get("tags") or [], ensure_ascii=False),
    }
    cp = getattr(doc, "custom_properties", None)
    if cp is None:
        return
    for k, v in payload.items():
        if v is None:
            continue
        _set_or_create_custom_property(cp, k, v)


def _write_docx(file_path: Path, fields: dict[str, Any]) -> None:
    from docx import Document
    doc = Document(str(file_path))
    _write_office_core(doc, fields)
    _write_office_custom(doc, fields)
    doc.save(str(file_path))


def _write_xlsx(file_path: Path, fields: dict[str, Any]) -> None:
    from openpyxl import load_workbook
    wb = load_workbook(str(file_path))
    cp = wb.properties
    title = _trim(fields.get("title"), 200)
    summary = _trim(fields.get("summary"), 500)
    tags = fields.get("tags") or []
    keywords = ", ".join(str(t) for t in tags) if tags else None
    category = fields.get("category")
    if title is not None:
        cp.title = title
    if summary is not None:
        cp.subject = summary
        cp.description = summary
    if keywords:
        cp.keywords = keywords
    if category:
        cp.category = str(category)
    wb.save(str(file_path))


def _write_pptx(file_path: Path, fields: dict[str, Any]) -> None:
    from pptx import Presentation
    pres = Presentation(str(file_path))
    cp = pres.core_properties
    title = _trim(fields.get("title"), 200)
    summary = _trim(fields.get("summary"), 500)
    tags = fields.get("tags") or []
    keywords = ", ".join(str(t) for t in tags) if tags else None
    category = fields.get("category")
    if title is not None:
        cp.title = title
    if summary is not None:
        cp.subject = summary
        cp.comments = summary
    if keywords:
        cp.keywords = keywords
    if category:
        cp.category = str(category)
    pres.save(str(file_path))


def _write_image(file_path: Path, fields: dict[str, Any]) -> None:
    """Write XMP packet + EXIF ImageDescription via Pillow."""
    try:
        from PIL import Image
    except ImportError as exc:
        raise NativeMetadataError("Pillow is not installed") from exc

    title = _trim(fields.get("title"), 200)
    summary = _trim(fields.get("summary"), 500)
    tags = fields.get("tags") or []
    category = fields.get("category")

    suffix = file_path.suffix.lower()

    with Image.open(file_path) as im:
        save_kwargs: dict[str, Any] = {}
        if suffix in {".jpg", ".jpeg"}:
            try:
                exif = im.getexif()
                if summary:
                    exif[0x010E] = summary
                if title:
                    exif[0x0131] = title
                save_kwargs["exif"] = exif.tobytes()
            except Exception:
                pass
            xmp = _build_xmp_packet(title=title, description=summary, tags=tags, category=category, correspondent=fields.get("correspondent"))
            save_kwargs["xmp"] = xmp.encode("utf-8")
            im.save(file_path, **save_kwargs)
        elif suffix == ".png":
            from PIL.PngImagePlugin import PngInfo
            info = PngInfo()
            if title:
                info.add_itxt("Title", title)
            if summary:
                info.add_itxt("Description", summary)
            if tags:
                info.add_itxt("Keywords", ", ".join(str(t) for t in tags))
            if category:
                info.add_itxt("Automafile:Category", str(category))
            xmp = _build_xmp_packet(title=title, description=summary, tags=tags, category=category, correspondent=fields.get("correspondent"))
            info.add_itxt("XML:com.adobe.xmp", xmp)
            im.save(file_path, pnginfo=info)
        elif suffix in {".tif", ".tiff"}:
            try:
                exif = im.getexif()
                if summary:
                    exif[0x010E] = summary
                if title:
                    exif[0x0131] = title
                save_kwargs["exif"] = exif
            except Exception:
                pass
            im.save(file_path, **save_kwargs)
        elif suffix == ".webp":
            xmp = _build_xmp_packet(title=title, description=summary, tags=tags, category=category, correspondent=fields.get("correspondent"))
            try:
                im.save(file_path, xmp=xmp.encode("utf-8"))
            except TypeError:
                im.save(file_path)
        else:
            raise NativeMetadataError(f"Unsupported image format: {suffix}")


def _xml_escape(s: str) -> str:
    return (
        s.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
        .replace("'", "&apos;")
    )


def _build_xmp_packet(
    *,
    title: str | None,
    description: str | None,
    tags: list[str] | None,
    category: str | None,
    correspondent: str | None,
) -> str:
    rdf_inner_parts: list[str] = []
    if title:
        rdf_inner_parts.append(
            f'      <dc:title><rdf:Alt><rdf:li xml:lang="x-default">{_xml_escape(title)}</rdf:li></rdf:Alt></dc:title>'
        )
    if description:
        rdf_inner_parts.append(
            f'      <dc:description><rdf:Alt><rdf:li xml:lang="x-default">{_xml_escape(description)}</rdf:li></rdf:Alt></dc:description>'
        )
    if tags:
        items = "\n".join(f"        <rdf:li>{_xml_escape(str(t))}</rdf:li>" for t in tags)
        rdf_inner_parts.append(
            f'      <dc:subject><rdf:Bag>\n{items}\n      </rdf:Bag></dc:subject>'
        )
    if category:
        rdf_inner_parts.append(
            f'      <{XMP_PREFIX}:Category>{_xml_escape(str(category))}</{XMP_PREFIX}:Category>'
        )
    if correspondent:
        rdf_inner_parts.append(
            f'      <{XMP_PREFIX}:Correspondent>{_xml_escape(str(correspondent))}</{XMP_PREFIX}:Correspondent>'
        )
    rdf_inner_parts.append(
        f'      <xmp:MetadataDate>{_utc_iso()}</xmp:MetadataDate>'
    )
    inner = "\n".join(rdf_inner_parts)
    return (
        '<?xpacket begin="﻿" id="W5M0MpCehiHzreSzNTczkc9d"?>\n'
        '<x:xmpmeta xmlns:x="adobe:ns:meta/">\n'
        '  <rdf:RDF xmlns:rdf="http://www.w3.org/1999/02/22-rdf-syntax-ns#">\n'
        '    <rdf:Description rdf:about="" '
        'xmlns:dc="http://purl.org/dc/elements/1.1/" '
        'xmlns:xmp="http://ns.adobe.com/xap/1.0/" '
        f'xmlns:{XMP_PREFIX}="{XMP_NAMESPACE}">\n'
        f'{inner}\n'
        '    </rdf:Description>\n'
        '  </rdf:RDF>\n'
        '</x:xmpmeta>\n'
        '<?xpacket end="w"?>'
    )


def read_native(file_path: Path) -> dict[str, Any]:
    """Best-effort read of native metadata (returns flat dict, never raises)."""
    suffix = file_path.suffix.lower()
    try:
        if suffix == ".pdf":
            from automafile.extractors.pdf import _read_pikepdf_metadata
            data, _ = _read_pikepdf_metadata(file_path)
            return data
        if suffix in {".docx", ".xlsx", ".pptx"}:
            from automafile.dispatch import EXT_MAP
            mod = EXT_MAP.get(suffix)
            if mod is None:
                return {}
            return mod.extract(file_path).native_metadata
        if suffix in {".jpg", ".jpeg", ".png", ".tiff", ".tif", ".webp", ".heic", ".heif"}:
            from automafile.extractors.image import _read_image_metadata
            return _read_image_metadata(file_path)
    except Exception:
        return {}
    return {}
