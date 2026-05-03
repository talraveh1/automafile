"""Verify each extractor populates ``ExtractedDoc.extracted_metadata`` with
the file's own embedded fields (Windows-11-style file properties)."""

from __future__ import annotations

from pathlib import Path

import pytest

from dragndoc.extractors._meta import clean_value, collect


# ---- helpers ----------------------------------------------------------


def test_clean_value_drops_empty_and_binary():
    assert clean_value(None) is None
    assert clean_value("") is None
    assert clean_value("   ") is None
    assert clean_value(b"\x00\x01\x02") is None
    assert clean_value("has \x00 nul") is None
    # XMP packet preamble — should be dropped wholesale, not clipped
    assert clean_value("<?xpacket begin='?'><x:xmpmeta>...") is None


def test_clean_value_clips_long_strings():
    long = "x" * 500
    out = clean_value(long)
    assert isinstance(out, str)
    assert len(out) <= 201  # 200 chars + ellipsis
    assert out.endswith("…")


def test_clean_value_keeps_numbers_and_lists():
    assert clean_value(42) == 42
    assert clean_value(3.14) == 3.14
    assert clean_value(False) is False  # explicit kept
    assert clean_value(["a", "b", "", None, "c"]) == ["a", "b", "c"]
    assert clean_value([]) is None


def test_collect_skips_known_binary_keys():
    raw = {
        "icc_profile": b"\x00" * 100,
        "XML:com.adobe.xmp": "<?xpacket...",
        "title": "Real Title",
    }
    out = collect(raw, prefix="info_")
    assert out == {"info_title": "Real Title"}


# ---- per-format extractors -------------------------------------------


def test_html_extractor_collects_title_and_meta(tmp_path):
    from dragndoc.extractors import html as html_ext
    p = tmp_path / "page.html"
    p.write_text(
        '<html lang="en"><head>'
        '<title>Quarterly Report</title>'
        '<meta name="author" content="Alice">'
        '<meta property="og:description" content="Q4 numbers">'
        '<meta http-equiv="content-language" content="en-US">'
        '</head><body>body text</body></html>',
        encoding="utf-8",
    )
    doc = html_ext.extract(p)
    md = doc.extracted_metadata
    assert md["title"] == "Quarterly Report"
    assert md["meta_author"] == "Alice"
    assert md["meta_og:description"] == "Q4 numbers"
    assert md["meta_content-language"] == "en-US"
    assert md["html_lang"] == "en"


def test_pdf_extractor_surfaces_docinfo(tmp_path):
    """DocInfo-only path: no XMP touch, so pikepdf doesn't sync-and-strip."""
    pikepdf = pytest.importorskip("pikepdf")
    from dragndoc.extractors import pdf as pdf_ext
    p = tmp_path / "doc.pdf"
    pdf = pikepdf.new()
    pdf.add_blank_page(page_size=(612, 792))
    pdf.docinfo["/Title"] = "Lease Agreement"
    pdf.docinfo["/Author"] = "Acme Realty"
    pdf.docinfo["/Subject"] = "Year 2026"
    pdf.docinfo["/Keywords"] = "lease,real-estate"
    pdf.save(p)

    doc = pdf_ext.extract(p)
    md = doc.extracted_metadata
    assert md["info_Title"] == "Lease Agreement"
    assert md["info_Author"] == "Acme Realty"
    assert md["info_Subject"] == "Year 2026"
    assert md["info_Keywords"] == "lease,real-estate"


def test_pdf_extractor_surfaces_xmp(tmp_path):
    """XMP-only path: opening the metadata block writes XMP and re-derives DocInfo."""
    pikepdf = pytest.importorskip("pikepdf")
    from dragndoc.extractors import pdf as pdf_ext
    p = tmp_path / "doc.pdf"
    pdf = pikepdf.new()
    pdf.add_blank_page(page_size=(612, 792))
    with pdf.open_metadata() as xmp:
        xmp["dc:title"] = "Lease Agreement"
        xmp["dc:creator"] = ["Acme Realty"]
    pdf.save(p)

    doc = pdf_ext.extract(p)
    md = doc.extracted_metadata
    assert md["dc:title"] == "Lease Agreement"
    # multi-valued XMP fields come through as lists
    assert md["dc:creator"] == ["Acme Realty"]
    # DocInfo derived by pikepdf from XMP
    assert md.get("info_Title") == "Lease Agreement"


def test_docx_extractor_surfaces_core_properties(tmp_path):
    docx_lib = pytest.importorskip("docx")
    from dragndoc.extractors import docx as docx_ext
    p = tmp_path / "doc.docx"
    d = docx_lib.Document()
    d.core_properties.title = "Project Brief"
    d.core_properties.author = "Bob"
    d.core_properties.subject = "Q1 plan"
    d.core_properties.keywords = "plan, q1, brief"
    d.core_properties.category = "Work"
    d.core_properties.last_modified_by = "Bob"
    d.add_paragraph("body content")
    d.save(str(p))

    doc = docx_ext.extract(p)
    md = doc.extracted_metadata
    assert md["core_title"] == "Project Brief"
    assert md["core_author"] == "Bob"
    assert md["core_subject"] == "Q1 plan"
    assert md["core_keywords"] == "plan, q1, brief"
    assert md["core_category"] == "Work"
    assert md["core_last_modified_by"] == "Bob"


def test_xlsx_extractor_surfaces_core_and_custom_props(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.packaging.custom import StringProperty
    from dragndoc.extractors import xlsx as xlsx_ext
    p = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "data"
    wb.properties.title = "Sales 2026"
    wb.properties.creator = "Carol"
    wb.properties.subject = "monthly numbers"
    wb.properties.keywords = "sales, q1"
    # custom prop
    wb.custom_doc_props.append(StringProperty(name="ProjectCode", value="X-42"))
    wb.save(str(p))

    doc = xlsx_ext.extract(p)
    md = doc.extracted_metadata
    assert md["core_title"] == "Sales 2026"
    assert md["core_creator"] == "Carol"
    assert md["core_subject"] == "monthly numbers"
    assert md["core_keywords"] == "sales, q1"
    assert md["custom_ProjectCode"] == "X-42"


def test_pptx_extractor_surfaces_core_properties(tmp_path):
    pptx_lib = pytest.importorskip("pptx")
    from dragndoc.extractors import pptx as pptx_ext
    p = tmp_path / "deck.pptx"
    pres = pptx_lib.Presentation()
    pres.core_properties.title = "Kickoff Deck"
    pres.core_properties.author = "Dave"
    pres.core_properties.keywords = "kickoff, q1"
    pres.save(str(p))

    doc = pptx_ext.extract(p)
    md = doc.extracted_metadata
    assert md["core_title"] == "Kickoff Deck"
    assert md["core_author"] == "Dave"
    assert md["core_keywords"] == "kickoff, q1"


def test_image_extractor_surfaces_exif_and_dimensions(tmp_path):
    PIL = pytest.importorskip("PIL")
    from PIL import Image
    from dragndoc.extractors import image as image_ext

    p = tmp_path / "shot.jpg"
    im = Image.new("RGB", (320, 240), color=(128, 0, 128))
    exif = im.getexif()
    # Standard EXIF tag IDs (we use the public PIL.ExifTags mapping).
    exif[0x010F] = "NIKON CORPORATION"      # Make
    exif[0x0110] = "NIKON D7000"            # Model
    exif[0x013B] = "Eve Photographer"       # Artist
    exif[0x010E] = "Sample portrait"        # ImageDescription
    im.save(p, exif=exif.tobytes())

    doc = image_ext.extract(p)
    md = doc.extracted_metadata
    assert md.get("dimensions") == "320x240"
    assert md.get("color_mode") == "RGB"
    assert md.get("pil_format") == "JPEG"
    assert md.get("exif_Make") == "NIKON CORPORATION"
    assert md.get("exif_Model") == "NIKON D7000"
    assert md.get("exif_Artist") == "Eve Photographer"
    assert md.get("exif_ImageDescription") == "Sample portrait"


def test_text_extractor_has_empty_extracted_metadata(tmp_path):
    from dragndoc.extractors import text as text_ext
    p = tmp_path / "note.txt"
    p.write_text("hello", encoding="utf-8")
    doc = text_ext.extract(p)
    assert doc.extracted_metadata == {}


# ---- integration with pipeline._hints_for -----------------------------


def test_hints_for_merges_extracted_metadata(tmp_path):
    from dragndoc.extractors.base import ExtractedDoc, Section
    from dragndoc.pipeline import _hints_for

    p = tmp_path / "x.pdf"
    p.write_bytes(b"%PDF-1.4\n")
    doc = ExtractedDoc(
        path=p,
        sections=[Section(label=None, text="...", index=0)],
        format="pdf",
        extracted_metadata={"info_Title": "Lease", "dc:creator": ["Acme"]},
    )
    hints = _hints_for(doc)
    assert hints["filename"] == "x.pdf"
    assert hints["format"] == "pdf"
    assert hints["info_Title"] == "Lease"
    assert hints["dc:creator"] == ["Acme"]
